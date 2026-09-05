from __future__ import annotations

import json
import os
import uuid
from pathlib import Path
from typing import Any

import matplotlib.pyplot as plt
import pandas as pd
from docx import Document

from langchain_core.tools import tool


OUTPUT_DIR = Path(
    os.getenv("AGENT_OUTPUT_DIR", "generated_reports")
).resolve()

OUTPUT_DIR.mkdir(
    parents=True,
    exist_ok=True
)


def _to_dataframe(data: Any) -> pd.DataFrame:
    """
    Convert SQL query results into a pandas DataFrame.

    Supports:
    - list[dict]
    - JSON string containing list[dict]
    - dict containing rows/data
    """

    if isinstance(data, str):
        data = json.loads(data)

    if isinstance(data, dict):

        if "rows" in data:
            data = data["rows"]

        elif "data" in data:
            data = data["data"]

        else:
            data = [data]

    if not isinstance(data, list):
        raise ValueError(
            "Query result must be a list of rows."
        )

    return pd.DataFrame(data)


@tool
def visualise_query_result(
    title: str,
    query_result: Any,
    x_column: str,
    y_column: str,
    chart_type: str = "bar",
) -> str:
    """
    Create a visualization from SQL query results.

    Use this tool when the user asks to:

    - visualize results
    - create a chart
    - compare values visually
    - show a graph
    - plot sales/revenue/orders

    Supported chart types:

    - bar
    - line
    - pie
    """

    try:

        df = _to_dataframe(query_result)

        if df.empty:
            return (
                "No data was available "
                "to visualize."
            )

        if x_column not in df.columns:
            return (
                f"Column '{x_column}' "
                "was not found."
            )

        if y_column not in df.columns:
            return (
                f"Column '{y_column}' "
                "was not found."
            )

        filename = (
            f"chart_{uuid.uuid4().hex[:8]}.png"
        )

        output_path = OUTPUT_DIR / filename

        plt.figure(figsize=(10, 6))

        if chart_type == "bar":

            plt.bar(
                df[x_column].astype(str),
                df[y_column],
            )

            plt.xlabel(x_column)
            plt.ylabel(y_column)

        elif chart_type == "line":

            plt.plot(
                df[x_column].astype(str),
                df[y_column],
                marker="o",
            )

            plt.xlabel(x_column)
            plt.ylabel(y_column)

        elif chart_type == "pie":

            plt.pie(
                df[y_column],
                labels=df[x_column].astype(str),
                autopct="%1.1f%%",
            )

        else:

            return (
                "Unsupported chart type. "
                "Use bar, line, or pie."
            )

        plt.title(title)

        if chart_type != "pie":
            plt.xticks(
                rotation=45,
                ha="right"
            )

        plt.tight_layout()

        plt.savefig(
            output_path,
            dpi=150,
            bbox_inches="tight",
        )

        plt.close()

        return json.dumps(
            {
                "success": True,
                "type": "visualization",
                "format": "png",
                "file": filename,
                "message": (
                    "Chart generated successfully: "
                    f"{filename}"
                ),
            }
        )

    except Exception as exc:

        return json.dumps(
            {
                "success": False,
                "error": str(exc),
            }
        )


@tool
def generate_report_document(
    title: str,
    query_result: Any,
    description: str = "",
) -> str:
    """
    Generate a DOCX report from SQL query results.

    Use this tool when the user asks for:

    - a report
    - a document
    - a downloadable summary
    - an analysis document
    - a comparison report
    """

    try:

        df = _to_dataframe(query_result)

        if df.empty:
            return (
                "No data was available "
                "to generate the document."
            )

        filename = (
            f"report_{uuid.uuid4().hex[:8]}.docx"
        )

        output_path = OUTPUT_DIR / filename

        document = Document()

        document.add_heading(
            title,
            level=1
        )

        if description:
            document.add_paragraph(
                description
            )

        document.add_paragraph(
            f"Records analyzed: {len(df)}"
        )

        document.add_heading(
            "Summary",
            level=2
        )

        numeric_columns = (
            df.select_dtypes(
                include="number"
            ).columns.tolist()
        )

        if numeric_columns:

            for column in numeric_columns:

                document.add_paragraph(
                    f"{column}: "
                    f"total={df[column].sum():,.2f}, "
                    f"average={df[column].mean():,.2f}"
                )

        document.add_heading(
            "Data",
            level=2
        )

        table = document.add_table(
            rows=1,
            cols=len(df.columns)
        )

        table.style = "Table Grid"

        # Header
        for i, column in enumerate(df.columns):

            table.rows[0].cells[i].text = str(
                column
            )

        # Data
        for _, row in df.iterrows():

            cells = table.add_row().cells

            for i, value in enumerate(row):

                cells[i].text = str(value)

        document.save(output_path)

        return json.dumps(
            {
                "success": True,
                "type": "document",
                "format": "docx",
                "file": filename,
                "message": (
                    "Report generated successfully: "
                    f"{filename}"
                ),
            }
        )

    except Exception as exc:

        return json.dumps(
            {
                "success": False,
                "error": str(exc),
            }
        )